"""
Módulo de geração de relatórios (CSV e Excel).
"""

import csv
import logging
from datetime import datetime
from pathlib import Path

from pesquisa_prontuario.models import ResultadoPesquisa

logger = logging.getLogger(__name__)

DIRETORIO_RELATORIOS = Path("relatorios")


def _garantir_diretorio():
    """Cria o diretório de relatórios se não existir."""
    DIRETORIO_RELATORIOS.mkdir(exist_ok=True)


def gerar_csv(resultados: list[ResultadoPesquisa], caminho: str = None) -> str:
    """
    Gera relatório em CSV a partir dos resultados da pesquisa.

    Args:
        resultados: Lista de ResultadoPesquisa
        caminho: Caminho do arquivo (opcional, gera automaticamente)

    Returns:
        Caminho do arquivo gerado
    """
    _garantir_diretorio()

    if not caminho:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        caminho = str(DIRETORIO_RELATORIOS / f"pesquisa_prontuario_{timestamp}.csv")

    cabecalhos = [
        "Nome Completo",
        "CPF",
        "Endereço",
        "Data do Atendimento",
        "Critério Atendido",
        "Detalhes",
    ]

    registros = []
    for resultado in resultados:
        for paciente in resultado.pacientes_encontrados:
            registros.append(paciente.to_dict())

    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=cabecalhos, delimiter=";")
        writer.writeheader()
        writer.writerows(registros)

    logger.info(f"CSV gerado: {caminho} ({len(registros)} registros)")
    return caminho


def gerar_excel(resultados: list[ResultadoPesquisa], caminho: str = None) -> str:
    """
    Gera relatório em Excel (.xlsx) a partir dos resultados da pesquisa.

    Args:
        resultados: Lista de ResultadoPesquisa
        caminho: Caminho do arquivo (opcional, gera automaticamente)

    Returns:
        Caminho do arquivo gerado
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        logger.warning("openpyxl não instalado. Instale com: pip install openpyxl")
        logger.info("Gerando CSV como alternativa...")
        return gerar_csv(resultados, caminho.replace(".xlsx", ".csv") if caminho else None)

    _garantir_diretorio()

    if not caminho:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        caminho = str(DIRETORIO_RELATORIOS / f"pesquisa_prontuario_{timestamp}.xlsx")

    wb = Workbook()

    # --- Aba: Resultados ---
    ws = wb.active
    ws.title = "Resultados"

    cabecalhos = [
        "Nome Completo",
        "CPF",
        "Endereço",
        "Data do Atendimento",
        "Critério Atendido",
        "Detalhes",
    ]

    # Estilo do cabeçalho
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
    borda = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col, cab in enumerate(cabecalhos, 1):
        celula = ws.cell(row=1, column=col, value=cab)
        celula.font = header_font
        celula.fill = header_fill
        celula.alignment = Alignment(horizontal="center")
        celula.border = borda

    # Dados
    linha = 2
    for resultado in resultados:
        for paciente in resultado.pacientes_encontrados:
            dados = paciente.to_dict()
            for col, cab in enumerate(cabecalhos, 1):
                celula = ws.cell(row=linha, column=col, value=dados.get(cab, ""))
                celula.border = borda
                celula.alignment = Alignment(wrap_text=True)
            linha += 1

    # Ajustar larguras
    larguras = [35, 18, 50, 18, 30, 40]
    for col, largura in enumerate(larguras, 1):
        ws.column_dimensions[chr(64 + col)].width = largura

    # Filtro automático
    if linha > 2:
        ws.auto_filter.ref = f"A1:F{linha - 1}"

    # --- Aba: Resumo ---
    ws_resumo = wb.create_sheet("Resumo")
    ws_resumo.cell(row=1, column=1, value="Data da Pesquisa").font = Font(bold=True)
    ws_resumo.cell(row=1, column=2, value="Total na Agenda").font = Font(bold=True)
    ws_resumo.cell(row=1, column=3, value="Pacientes com Critérios").font = Font(bold=True)
    ws_resumo.cell(row=1, column=4, value="Erros").font = Font(bold=True)

    for i, resultado in enumerate(resultados, 2):
        ws_resumo.cell(
            row=i, column=1, value=resultado.data_pesquisa.strftime("%d/%m/%Y")
        )
        ws_resumo.cell(row=i, column=2, value=resultado.total_pacientes_agenda)
        ws_resumo.cell(row=i, column=3, value=resultado.total_encontrados)
        ws_resumo.cell(row=i, column=4, value=len(resultado.erros))

    ws_resumo.column_dimensions["A"].width = 18
    ws_resumo.column_dimensions["B"].width = 18
    ws_resumo.column_dimensions["C"].width = 25
    ws_resumo.column_dimensions["D"].width = 10

    wb.save(caminho)
    logger.info(f"Excel gerado: {caminho}")
    return caminho


def imprimir_resumo(resultados: list[ResultadoPesquisa]) -> None:
    """Imprime um resumo no terminal."""
    total_pacientes = 0
    total_encontrados = 0
    total_erros = 0

    print("\n" + "=" * 60)
    print("  RESUMO DA PESQUISA DE PRONTUÁRIOS")
    print("=" * 60)

    for resultado in resultados:
        total_pacientes += resultado.total_pacientes_agenda
        total_encontrados += resultado.total_encontrados
        total_erros += len(resultado.erros)

        print(f"\n  Data: {resultado.data_pesquisa.strftime('%d/%m/%Y')}")
        print(f"  Pacientes na agenda: {resultado.total_pacientes_agenda}")
        print(f"  Com critérios atendidos: {resultado.total_encontrados}")

        for pac in resultado.pacientes_encontrados:
            print(f"    - {pac.nome_completo} | {pac.criterio.value}")

        if resultado.erros:
            print(f"  Erros: {len(resultado.erros)}")

    print(f"\n{'─' * 60}")
    print(f"  TOTAIS:")
    print(f"    Pacientes verificados: {total_pacientes}")
    print(f"    Pacientes encontrados: {total_encontrados}")
    if total_erros:
        print(f"    Erros: {total_erros}")
    print("=" * 60 + "\n")
